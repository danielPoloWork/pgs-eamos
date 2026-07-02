#!/usr/bin/env python3
"""EAMOS emit_xlsx: data-IR -> a .xlsx KPI table.

The cosmetic IR -> binary hop for the data-IR (RFC-0002 §3), non-deterministic / non-dependency-free:
requires openpyxl (`pip install openpyxl`); the core never imports it. The KPI values are sourced
ledger cells (no formulas → zero formula errors by construction); each row carries its source as the
hardcode-provenance note. Assumed values get a **yellow highlight** — the spreadsheet-native idiom
for "needs attention" — the data-IR's form of the grounding signal (RFC-0001 §6).

    python tools/render.py manifest.yaml --ir data --out build/data-ir.json
    python tools/emit_xlsx.py build/data-ir.json --out build/kpi.xlsx
"""

import argparse
import json
import os
import sys

import _cli    # utf8_stdio (#54)
import labels  # chrome labels as data (#61)

SUPPORTED_IR_VERSION = 1   # the IR contract this emitter was written for (#62)

def _lab(lang, key):
    return labels.lab(lang, "xlsx", key)   # chrome labels as data (#61)


def build_xlsx(ir, out_path):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment

    lang = ir.get("output_lang", "en")
    navy = "1E2761"
    header_fill = PatternFill("solid", fgColor="EEF1F7")
    assumed_fill = PatternFill("solid", fgColor="FFF2CC")   # attention highlight (xlsx idiom)
    base = Font(name="Arial", size=11)
    bold = Font(name="Arial", size=11, bold=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "KPI"

    ws["A1"] = ir.get("title", "")
    ws["A1"].font = Font(name="Arial", size=14, bold=True, color=navy)
    ws.append([])

    headers = [_lab(lang, k) for k in ("kpi", "value", "target", "source")]
    ws.append(headers)
    hdr_row = ws.max_row
    for col in range(1, len(headers) + 1):
        c = ws.cell(row=hdr_row, column=col)
        c.font = bold
        c.fill = header_fill

    for r in ir.get("rows", []):
        note = r.get("source", "")
        ws.append([r.get("label", ""), r.get("value", ""), r.get("target", ""), note])
        row = ws.max_row
        for col in range(1, 5):
            ws.cell(row=row, column=col).font = base
        if r.get("assumed"):
            ws.cell(row=row, column=2).fill = assumed_fill   # the value cell needs attention

    appendix = ir.get("review_appendix", [])
    if appendix:
        ws.append([])
        ws.append(["⚠ " + _lab(lang, "review")])
        ws.cell(row=ws.max_row, column=1).font = Font(name="Arial", size=11, bold=True, color="B86B00")
        for a in appendix:
            line = f"{a.get('binding', '')} = {a.get('value', '')}"
            if a.get("assumption"):
                line += f" — {a['assumption']}"
            if a.get("fill_from"):
                line += f" ({a['fill_from']})"
            ws.append([line])
            ws.cell(row=ws.max_row, column=1).font = base

    for col, width in (("A", 34), ("B", 18), ("C", 16), ("D", 48)):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = ws.cell(row=hdr_row + 1, column=1)
    wb.properties.creator = "EAMOS"
    wb.save(out_path)
    return ws.max_row


def main():
    _cli.utf8_stdio()   # Windows: piped stdout must stay UTF-8 (#54)
    ap = argparse.ArgumentParser(description="Emit a .xlsx KPI table from a data-IR.")
    ap.add_argument("data_ir", help="path to a data-IR JSON file")
    ap.add_argument("--out", required=True, help="output .xlsx path")
    args = ap.parse_args()

    try:
        import openpyxl  # noqa: F401
    except ImportError:
        print("emit_xlsx: openpyxl is not installed. This is the cosmetic hop "
              "(outside the dependency-free core). Install it with: pip install openpyxl")
        return 2

    with open(args.data_ir, encoding="utf-8") as fh:
        ir = json.load(fh)
    _cli.require_ir_version(ir, "emit_xlsx", SUPPORTED_IR_VERSION)
    os.makedirs(os.path.dirname(os.path.abspath(args.out)), exist_ok=True)
    n = build_xlsx(ir, args.out)
    print(f"emit_xlsx: OK — KPI table ({n} rows) -> {args.out}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
