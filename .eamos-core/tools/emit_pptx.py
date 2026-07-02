#!/usr/bin/env python3
"""EAMOS emit_pptx: deck-IR -> a .pptx board deck.

The cosmetic IR -> binary hop (RFC-0001 §14): the ONLY non-deterministic / non-dependency-free
layer. Requires python-pptx (`pip install python-pptx`); the deterministic core (render.py,
eamos_lint.py) never imports it. The deck-IR is the determinism boundary — gates run there, not here.

    python tools/emit_pptx.py build/deck-ir.json --out build/deck.pptx

Assumed values (RFC-0001 §6) carry the ⟨… — da verificare⟩ marker from the IR and are rendered in
amber so the "verify before the room" items are visible at a glance.
"""

import argparse
import json
import os
import sys

import _cli  # utf8_stdio (#54)

NAVY = (0x1E, 0x27, 0x61)      # title
BODY = (0x2B, 0x2B, 0x2B)      # body text
AMBER = (0xB8, 0x6B, 0x00)     # assumed / to-verify
MUTED = (0x70, 0x70, 0x70)     # captions

LABELS = {
    "it": {"target": "target", "risk": "Rischio", "ask": "Richiesta", "pro": "pro", "con": "contro",
           "review": "Da verificare prima della sala", "fill": "fonte",
           "next_step": "Prossimo passo", "residual_risk": "Rischio residuo"},
    "en": {"target": "target", "risk": "Risk", "ask": "Ask", "pro": "pro", "con": "con",
           "review": "Verify before the room", "fill": "source",
           "next_step": "Next step", "residual_risk": "Residual risk"},
}


def _lab(lang, key):
    return LABELS.get(lang, LABELS["en"]).get(key, LABELS["en"][key])


def _line(block, lang):
    """Flatten one deck-IR block into a single display line of text."""
    t = block.get("type")
    if t in ("lead", "bullet", "prose", "decision"):
        return block.get("text", "")
    if t == "kpi_row":
        s = f"{block.get('label', '')}: {block.get('value', '')}"
        if "target" in block:
            s += f"  ({_lab(lang, 'target')}: {block['target']})"
        return s
    if t == "risk":
        return f"{_lab(lang, 'risk')}: {block.get('risk', '')} — {_lab(lang, 'ask')}: {block.get('ask', '')}"
    if t == "option":
        return f"{block.get('name', '')} — {_lab(lang, 'pro')}: {block.get('pro', '')}; " \
               f"{_lab(lang, 'con')}: {block.get('con', '')}"
    if t == "next_step":
        return f"{_lab(lang, 'next_step')}: {block.get('text', '')}"
    if t == "residual_risk":
        return f"{_lab(lang, 'residual_risk')}: {block.get('text', '')}"
    return ""


def build_pptx(ir, out_path):
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN

    lang = ir.get("output_lang", "en")
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    def textbox(slide, left, top, width, height):
        tf = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height)).text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = Inches(0)
        return tf

    def para(tf, text, size, color, bold=False, bullet=False, first=False):
        p = tf.paragraphs[0] if first and not tf.paragraphs[0].runs else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        run = p.add_run()
        run.text = ("•  " + text) if bullet else text
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.name = "Calibri"
        # Assumed/to-verify values carry the ⟨…⟩ marker from the IR — flag them in amber.
        run.font.color.rgb = RGBColor(*(AMBER if "⟨" in text else color))
        p.space_after = Pt(6)
        return p

    # Title slide.
    s0 = prs.slides.add_slide(blank)
    t = textbox(s0, 0.7, 2.4, 12.0, 2.4)
    para(t, ir.get("objective", ""), 40, NAVY, bold=True, first=True)
    cap = textbox(s0, 0.7, 4.7, 12.0, 0.6)
    para(cap, f"{ir.get('archetype', '')}  ·  {ir.get('altitude', '')}  ·  "
              f"{ir.get('deliverable', '')}/{ir.get('format', '')}  ·  {lang}", 16, MUTED, first=True)

    # Content slides.
    for slide in ir.get("slides", []):
        sl = prs.slides.add_slide(blank)
        head = textbox(sl, 0.7, 0.5, 12.0, 1.0)
        para(head, slide.get("title", slide.get("id", "")), 32, NAVY, bold=True, first=True)
        body = textbox(sl, 0.7, 1.7, 12.0, 5.3)
        for i, b in enumerate(slide.get("blocks", [])):
            line = _line(b, lang)
            if not line:
                continue
            if b.get("type") == "lead":
                para(body, line, 18, BODY, bold=True, first=(i == 0))
            elif b.get("type") == "kpi_row":
                para(body, line, 16, BODY, bullet=True, first=(i == 0))
            else:
                para(body, line, 16, BODY, bullet=True, first=(i == 0))

    # Review appendix slide.
    appendix = ir.get("review_appendix", [])
    if appendix:
        sl = prs.slides.add_slide(blank)
        head = textbox(sl, 0.7, 0.5, 12.0, 1.0)
        para(head, "⚠ " + _lab(lang, "review"), 32, AMBER, bold=True, first=True)
        body = textbox(sl, 0.7, 1.7, 12.0, 5.3)
        for i, a in enumerate(appendix):
            line = f"{a.get('binding', '')} = {a.get('value', '')}"
            if a.get("assumption"):
                line += f" — {a['assumption']}"
            if a.get("fill_from"):
                line += f" ({_lab(lang, 'fill')}: {a['fill_from']})"
            para(body, line, 15, AMBER, bullet=True, first=(i == 0))

    prs.core_properties.author = "EAMOS"
    prs.core_properties.title = ir.get("objective", "")[:120]
    prs.save(out_path)
    return len(prs.slides._sldIdLst)


def main():
    _cli.utf8_stdio()   # Windows: piped stdout must stay UTF-8 (#54)
    ap = argparse.ArgumentParser(description="Emit a .pptx board deck from a deck-IR.")
    ap.add_argument("deck_ir", help="path to a deck-IR JSON file")
    ap.add_argument("--out", required=True, help="output .pptx path")
    args = ap.parse_args()

    try:
        import pptx  # noqa: F401
    except ImportError:
        print("emit_pptx: python-pptx is not installed. This is the cosmetic hop "
              "(outside the dependency-free core). Install it with: pip install python-pptx")
        return 2

    with open(args.deck_ir, encoding="utf-8") as fh:
        ir = json.load(fh)
    os.makedirs(os.path.dirname(os.path.abspath(args.out)), exist_ok=True)
    n = build_pptx(ir, args.out)
    print(f"emit_pptx: OK — {n} slides -> {args.out}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
