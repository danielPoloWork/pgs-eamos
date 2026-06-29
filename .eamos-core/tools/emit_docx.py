#!/usr/bin/env python3
"""EAMOS emit_docx: deck-IR -> a .docx pre-read.

The cosmetic IR -> binary hop (RFC-0001 §14), non-deterministic / non-dependency-free: requires
python-docx (`pip install python-docx`); the core never imports it. Unlike the deck (slides), the
pre-read is the *read-alone* prose projection of the same grounded content (RFC-0002 §3, doc-IR):
headings, prose, bullets, a KPI table, and the "verify before the room" appendix.

    python tools/emit_docx.py build/deck-ir.json --out build/pre-read.docx

Assumed values keep their ⟨… — da verificare⟩ marker from the IR and are rendered in amber.
"""

import argparse
import json
import os
import sys

AMBER = (0xB8, 0x6B, 0x00)
MUTED = (0x70, 0x70, 0x70)

LABELS = {
    "it": {"target": "Target", "value": "Valore", "kpi": "KPI", "risk": "Rischio", "ask": "Richiesta",
           "pro": "pro", "con": "contro", "review": "Da verificare prima della sala",
           "fill": "fonte", "preread": "Pre-read",
           "next_step": "Prossimo passo", "residual_risk": "Rischio residuo"},
    "en": {"target": "Target", "value": "Value", "kpi": "KPI", "risk": "Risk", "ask": "Ask",
           "pro": "pro", "con": "con", "review": "Verify before the room",
           "fill": "source", "preread": "Pre-read",
           "next_step": "Next step", "residual_risk": "Residual risk"},
}


def _lab(lang, key):
    return LABELS.get(lang, LABELS["en"]).get(key, LABELS["en"][key])


def _run(paragraph, text, bold=False, color=None):
    from docx.shared import RGBColor
    run = paragraph.add_run(text)
    run.bold = bold
    rgb = AMBER if (color is None and "⟨" in (text or "")) else color
    if rgb:
        run.font.color.rgb = RGBColor(*rgb)
    return run


def build_docx(ir, out_path):
    from docx import Document
    from docx.shared import Pt

    lang = ir.get("output_lang", "en")
    doc = Document()

    doc.add_heading(ir.get("objective", ""), level=0)
    cap = doc.add_paragraph()
    _run(cap, f"{_lab(lang, 'preread')} · {ir.get('archetype', '')} · {ir.get('altitude', '')} · "
              f"{ir.get('deliverable', '')}/{ir.get('format', '')} · {lang}", color=MUTED)

    for slide in ir.get("slides", []):
        doc.add_heading(slide.get("title", slide.get("id", "")), level=1)
        blocks = slide.get("blocks", [])
        if slide.get("kind") == "kpi_table":
            table = doc.add_table(rows=1, cols=3)
            table.style = "Table Grid"
            hdr = table.rows[0].cells
            for cell, key in zip(hdr, ("kpi", "value", "target")):
                _run(cell.paragraphs[0], _lab(lang, key), bold=True)
            for b in blocks:
                cells = table.add_row().cells
                _run(cells[0].paragraphs[0], b.get("label", ""))
                _run(cells[1].paragraphs[0], b.get("value", ""))
                _run(cells[2].paragraphs[0], b.get("target", ""))
            continue
        for b in blocks:
            t = b.get("type")
            if t == "lead":
                _run(doc.add_paragraph(), b.get("text", ""), bold=True)
            elif t == "prose":
                _run(doc.add_paragraph(), b.get("text", ""))
            elif t == "bullet":
                _run(doc.add_paragraph(style="List Bullet"), b.get("text", ""))
            elif t == "decision":
                _run(doc.add_paragraph(style="List Number"), b.get("text", ""))
            elif t == "risk":
                p = doc.add_paragraph(style="List Bullet")
                _run(p, f"{_lab(lang, 'risk')}: ", bold=True)
                _run(p, b.get("risk", ""))
                _run(p, f" — {_lab(lang, 'ask')}: ", bold=True)
                _run(p, b.get("ask", ""))
            elif t == "option":
                p = doc.add_paragraph(style="List Bullet")
                _run(p, b.get("name", ""), bold=True)
                _run(p, f" — {_lab(lang, 'pro')}: {b.get('pro', '')}; "
                        f"{_lab(lang, 'con')}: {b.get('con', '')}")
            elif t == "next_step":
                p = doc.add_paragraph()
                _run(p, f"{_lab(lang, 'next_step')}: ", bold=True)
                _run(p, b.get("text", ""))
            elif t == "residual_risk":
                p = doc.add_paragraph(style="List Bullet")
                _run(p, f"{_lab(lang, 'residual_risk')}: ", bold=True)
                _run(p, b.get("text", ""))

    appendix = ir.get("review_appendix", [])
    if appendix:
        doc.add_heading("⚠ " + _lab(lang, "review"), level=1)
        for a in appendix:
            p = doc.add_paragraph(style="List Bullet")
            line = f"{a.get('binding', '')} = {a.get('value', '')}"
            if a.get("assumption"):
                line += f" — {a['assumption']}"
            if a.get("fill_from"):
                line += f" ({_lab(lang, 'fill')}: {a['fill_from']})"
            _run(p, line, color=AMBER)

    doc.core_properties.author = "EAMOS"
    doc.core_properties.title = (ir.get("objective", "") or "")[:120]
    doc.save(out_path)
    return len(doc.paragraphs)


def main():
    ap = argparse.ArgumentParser(description="Emit a .docx pre-read from a deck-IR.")
    ap.add_argument("deck_ir", help="path to a deck-IR JSON file")
    ap.add_argument("--out", required=True, help="output .docx path")
    args = ap.parse_args()

    try:
        import docx  # noqa: F401
    except ImportError:
        print("emit_docx: python-docx is not installed. This is the cosmetic hop "
              "(outside the dependency-free core). Install it with: pip install python-docx")
        return 2

    with open(args.deck_ir, encoding="utf-8") as fh:
        ir = json.load(fh)
    os.makedirs(os.path.dirname(os.path.abspath(args.out)), exist_ok=True)
    n = build_docx(ir, args.out)
    print(f"emit_docx: OK — pre-read ({n} paragraphs) -> {args.out}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
